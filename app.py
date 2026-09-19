import os
import re
import io
import json
import shutil
import mimetypes
from datetime import datetime
import urllib.parse
import pandas as pd
import openpyxl
from PIL import Image, ImageOps, ImageFile
import gradio as gr
from fastapi import FastAPI, Response
from fastapi.responses import FileResponse
from pydantic import BaseModel

ImageFile.LOAD_TRUNCATED_IMAGES = True

# ==============================
# 0. 本地配置持久化
# ==============================
CONFIG_DIR = os.path.join(os.getcwd(), "gradio_temp")
CONFIG_FILE = os.path.join(CONFIG_DIR, "config.json")

def load_config():
    if os.path.exists(CONFIG_FILE):
        try:
            with open(CONFIG_FILE, 'r', encoding='utf-8') as f:
                return json.load(f)
        except:
            pass
    return {"path_a": "", "path_b": "", "asset_dir": "", "video_dir": ""}

def save_config(path_a, path_b, asset_dir, video_dir):
    os.makedirs(CONFIG_DIR, exist_ok=True)
    with open(CONFIG_FILE, 'w', encoding='utf-8') as f:
        json.dump({
            "path_a": path_a,
            "path_b": path_b,
            "asset_dir": asset_dir,
            "video_dir": video_dir
        }, f, ensure_ascii=False, indent=2)

# ==============================
# 1. 资产索引与工具函数
# ==============================
def clean_path(path_str: str) -> str:
    if not path_str:
        return ""
    p = path_str.strip().strip("'").strip('"')
    p = p.replace("\\ ", " ")
    return os.path.abspath(p)

def build_image_index(asset_dir: str):
    index = {}
    if not asset_dir or not os.path.isdir(asset_dir):
        return index
    valid_exts = {'.jpg', '.jpeg', '.png', '.webp', '.bmp', '.tif', '.tiff'}
    pattern = re.compile(r"^([a-zA-Z0-9]+)[._\-]")
    for root, _, files in os.walk(asset_dir):
        for f in files:
            ext = os.path.splitext(f)[1].lower()
            if ext in valid_exts:
                match = pattern.match(f)
                if match:
                    code = match.group(1).upper()
                else:
                    code = os.path.splitext(f)[0].upper()
                index[code] = os.path.join(root, f)
    return index

def build_audio_index(asset_dir: str):
    index = {}
    if not asset_dir or not os.path.isdir(asset_dir):
        return index
    valid_exts = {'.mp3', '.wav', '.aac', '.m4a', '.flac'}
    pattern = re.compile(r"^([a-zA-Z0-9]+)[._\-]")
    for root, _, files in os.walk(asset_dir):
        for f in files:
            ext = os.path.splitext(f)[1].lower()
            if ext in valid_exts:
                match = pattern.match(f)
                if match:
                    code = match.group(1).upper()
                else:
                    code = os.path.splitext(f)[0].upper()
                index[code] = os.path.join(root, f)
    return index

def build_video_list(video_dir: str):
    video_paths = []
    if not video_dir or not os.path.isdir(video_dir):
        return video_paths
    video_exts = {'.mp4', '.mov', '.webm', '.mkv'}
    for root, _, files in os.walk(video_dir):
        for f in files:
            ext = os.path.splitext(f)[1].lower()
            if ext in video_exts:
                video_paths.append(os.path.join(root, f))
    return video_paths

def find_matching_video(shot_id: str, video_paths: list):
    shot_lower = shot_id.lower()
    separators = {'.', '-', '_'}
    for path in video_paths:
        fname = os.path.basename(path)
        name_body, _ = os.path.splitext(fname)
        name_lower = name_body.lower()
        if name_lower == shot_lower:
            return path
        if any(name_lower.startswith(f"{shot_lower}{sep}") for sep in separators):
            return path
    return None

def backup_excel(filepath: str):
    if not os.path.exists(filepath):
        return
    backup_dir = os.path.join(os.path.dirname(filepath), "_backups")
    os.makedirs(backup_dir, exist_ok=True)
    basename = os.path.basename(filepath)
    name, ext = os.path.splitext(basename)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = os.path.join(backup_dir, f"{name}_{timestamp}{ext}")
    shutil.copy2(filepath, backup_path)

def natural_sort_key(s: str):
    return [int(text) if text.isdigit() else text.lower() for text in re.split(r'(\d+)', str(s))]

# ==============================
# 2. 核心表格渲染引擎
# ==============================
def render_table_b_html(df, image_index, audio_index, video_paths, sheet_name):
    if df is None or df.empty:
        return "<div class='empty-tip'>暂无数据或工作表为空</div>"

    rows_html = []

    for idx, row in df.iterrows():
        shot_id = str(row.iloc[0]) if len(row) > 0 and pd.notna(row.iloc[0]) else f"Sht{idx+1:02d}"
        title = str(row.iloc[1]) if len(row) > 1 and pd.notna(row.iloc[1]) else ""
        try:
            raw_dur = float(row.iloc[2]) if len(row) > 2 and pd.notna(row.iloc[2]) else 0.0
        except:
            raw_dur = 0.0
        img_codes_str = str(row.iloc[3]) if len(row) > 3 and pd.notna(row.iloc[3]) else ""
        audio_codes_str = str(row.iloc[4]) if len(row) > 4 and pd.notna(row.iloc[4]) else ""
        prompt = str(row.iloc[5]) if len(row) > 5 and pd.notna(row.iloc[5]) else ""

        vid_path = find_matching_video(shot_id, video_paths)

        # 图片资产
        img_items = [c.strip() for c in img_codes_str.split('|') if c.strip() and c.strip() != '0']
        img_previews = []
        for code in img_items[:6]:
            code_upper = code.upper()
            if code_upper in image_index:
                file_url = f"/api/local_media?filepath={urllib.parse.quote(image_index[code_upper])}"
                img_previews.append(f"""
                    <div class="asset-card" onclick="openLightbox('{file_url}')" title="点击放大: {code}">
                        <img src="{file_url}" alt="{code}" class="thumbnail-img"/>
                        <span class="asset-badge">{code}</span>
                    </div>""")
            else:
                img_previews.append(f"""<div class="asset-card missing" title="未找到文件"><span class="asset-badge">{code}</span><div class="missing-placeholder">无文件</div></div>""")

        img_grid = f"<div class='img-grid'>{''.join(img_previews)}</div>" if img_previews else ""
        img_cell_content = f"""<div class="cell-flex-wrapper">{img_grid}<input type="text" class="raw-text-edit" data-col="img" value="{img_codes_str}" placeholder="图片资产"/></div>"""

        # 声音资产
        audio_items = [c.strip() for c in audio_codes_str.split('|') if c.strip() and c.strip() != '0']
        audio_previews = []
        for code in audio_items[:3]:
            code_upper = code.upper()
            if code_upper in audio_index:
                file_url = f"/api/local_media?filepath={urllib.parse.quote(audio_index[code_upper])}"
                audio_previews.append(f"""
                <div class="audio-card">
                    <span class="asset-badge-audio">{code}</span>
                    <audio controls preload="auto" src="{file_url}" class="audio-player" onclick="event.stopPropagation()">
                </div>""")
            else:
                audio_previews.append(f"""
                <div class="audio-card missing">
                    <span class="asset-badge-audio">{code} (未找到)</span>
                </div>""")

        audio_stack = f"<div class='audio-stack'>{''.join(audio_previews)}</div>" if audio_previews else ""
        audio_cell_content = f"""<div class="cell-flex-wrapper">{audio_stack}<input type="text" class="raw-text-edit" data-col="audio" value="{audio_codes_str}" placeholder="声音资产"/></div>"""

        # 视频结果：默认无控件，悬停显示
        video_preview = ""
        if vid_path:
            vid_url = f"/api/local_media?filepath={urllib.parse.quote(vid_path)}"
            video_preview = f"""
            <div class="video-cell-wrapper">
                <div class="video-wrap" 
                     onmouseenter="this.querySelector('video').controls = true"
                     onmouseleave="this.querySelector('video').controls = false">
                    <video class="video-player"
                           preload="metadata"
                           onplay="pauseOtherVideos(this)">
                        <source src="{vid_url}">
                    </video>
                    <button class="video-expand-btn" onclick="openVideoLightbox('{vid_url}', this)" title="放大查看">⛶</button>
                </div>
                <input type="text" class="raw-text-edit" data-col="video" value="" placeholder="自动匹配"/>
            </div>
            """

        video_cell_content = f"""<div class="cell-flex-wrapper">{video_preview}</div>"""

        rows_html.append(f"""
            <tr data-row-idx="{idx}" class="storyboard-row" onclick="onRowClick(this)">
                <td class="col-shot"><input type="text" class="raw-text-edit cell-center shot-input" data-col="shot" value="{shot_id}"/></td>
                <td class="col-title"><textarea class="raw-text-edit area-title" data-col="title">{title}</textarea></td>
                <td class="col-dur">
                    <input type="text" class="raw-text-edit cell-center dur-input" data-col="dur" value="{raw_dur}"/>
                </td>
                <td class="col-img">{img_cell_content}</td>
                <td class="col-audio">{audio_cell_content}</td>
                <td class="col-prompt"><div class="prompt-box" data-col="prompt" contenteditable="true">{prompt}</div></td>
                <td class="col-video">{video_cell_content}</td>
            </tr>
        """)

    table_body = "".join(rows_html)

    return f"""
    <div class="storyboard-viewport" id="storyboardViewport" data-sheet-name="{sheet_name}">
        <table class="storyboard-table" id="storyboardTable">
            <thead>
                <tr>
                    <th class="col-shot">A.镜头号</th>
                    <th class="col-title">B.小标题</th>
                    <th class="col-dur">C.时长</th>
                    <th class="col-img">D.图片资产 (预览+编辑)</th>
                    <th class="col-audio">E.声音资产 (预览+编辑)</th>
                    <th class="col-prompt">F.生视频提示词 (点击本行大字显示)</th>
                    <th class="col-video">G.视频结果 (预览+编辑)</th>
                </tr>
            </thead>
            <tbody>{table_body}</tbody>
        </table>
    </div>
    """

# ==============================
# 3. CSS 样式
# ==============================
CUSTOM_CSS = """
:root {
    --col-scale: 1.0;
    --row-scale: 1.0;
    --img-base-w: 80px;
    --img-base-h: 56px;
    --header-height: 44px;
}

.hidden-element { display: none !important; }

.gradio-container {
    max-width: 98% !important;
    background-color: #121418 !important;
    color: #e1e4ea !important;
}

/* Toast 提示 */
.toast-box {
    position: fixed;
    top: 20px;
    left: 50%;
    transform: translateX(-50%);
    z-index: 10000;
    padding: 12px 24px;
    border-radius: 8px;
    font-size: 14px;
    font-weight: 500;
    box-shadow: 0 4px 12px rgba(0,0,0,0.3);
    opacity: 0;
    transition: opacity 0.3s, top 0.3s;
    pointer-events: none;
}
.toast-success { background: #2d5a3d; color: #d4edda; border: 1px solid #3d7a4f; }
.toast-error { background: #5c2e2e; color: #f8d7da; border: 1px solid #7a3a3a; }
.toast-show { opacity: 1; top: 30px; }

/* 顶部工具栏 第一行 */
.top-toolbar-row {
    display: flex;
    align-items: center;
    gap: 10px;
    padding: 8px 10px;
    background: #1c2027;
    border: 1px solid #2b303c;
    border-radius: 8px;
    margin-bottom: 8px;
}

/* 复选框更紧凑 */
.col-checkbox-wrap {
    flex: 1;
    display: flex;
    align-items: center;
}

/* 顶部工具栏 第二行：视图微调 */
.view-toolbar-row {
    display: flex;
    align-items: center;
    gap: 8px;
    padding: 6px 10px;
    background: #1c2027;
    border: 1px solid #2b303c;
    border-radius: 8px;
    margin-bottom: 8px;
    flex-wrap: nowrap;
}

.view-toolbar-row button {
    min-height: 30px !important;
    font-size: 12px !important;
    padding: 0 10px !important;
    flex-shrink: 0;
}

/* 主提示区 + 保存按钮行 */
.prompt-save-row {
    display: flex;
    align-items: stretch;
    gap: 10px;
    margin-bottom: 8px;
}

.master-prompt-preview {
    flex: 1;
    padding: 10px 12px;
    background: #1c2027;
    border-left: 4px solid #ebcb8b;
    border-radius: 4px;
    height: 122px;
    box-sizing: border-box;
    overflow-y: auto;
    overflow-x: hidden;
    font-size: 15px;
    line-height: 1.55;
    color: #f0f3f8;
    white-space: pre-wrap !important;
    word-wrap: break-word !important;
    overflow-wrap: break-word !important;
    word-break: break-word;
}

.master-prompt-preview:empty::before {
    content: "👈 点击下方表格任意行，此处将大字号展示该镜头的完整生视频提示词。";
    color: #8c9ba5;
}

.save-btn {
    min-height: 30px !important;
    align-self: stretch;
    min-width: 180px !important;
}

/* 表格视口 */
.storyboard-viewport {
    width: 100%;
    overflow: auto;
    max-height: 70vh;
    border: 1px solid #2b303c;
    border-radius: 8px;
    background: #181a20;
    scroll-behavior: smooth;
    -webkit-overflow-scrolling: touch;
}

.storyboard-table {
    width: max-content;
    min-width: 100%;
    border-collapse: collapse;
    font-size: 13px;
    table-layout: fixed;
}

.storyboard-table thead th {
    position: sticky;
    top: 0;
    background: #20242c;
    color: #c8d1e0;
    font-weight: 600;
    border-bottom: 2px solid #2e3440;
    z-index: 10;
    text-align: left;
    height: var(--header-height);
    line-height: var(--header-height);
    padding: 0 8px;
    white-space: nowrap;
}

.storyboard-table td {
    border-bottom: 1px solid #282c37;
    padding: 6px;
    vertical-align: top;
}

.storyboard-row:hover {
    background-color: #1a1e27;
    cursor: pointer;
}

.storyboard-row.selected-row {
    background-color: #242e3f !important;
    box-shadow: inset 4px 0 0 #88c0d0;
}

/* 列宽 */
.col-shot { width: calc(78px * var(--col-scale)); }
.col-title { width: calc(145px * var(--col-scale)); }
.col-dur { width: calc(56px * var(--col-scale)); }
.col-img { width: calc(300px * var(--col-scale)); }
.col-audio { width: calc(290px * var(--col-scale)); }
.col-prompt { width: calc(360px * var(--col-scale)); }
.col-video { width: calc(240px * var(--col-scale)); }

/* 图片网格 */
.img-grid {
    display: grid;
    grid-template-columns: repeat(3, calc(var(--img-base-w) * var(--col-scale)));
    gap: 5px;
    width: max-content;
}

.asset-card {
    position: relative;
    width: calc(var(--img-base-w) * var(--col-scale));
    height: calc(var(--img-base-h) * var(--col-scale));
    border-radius: 4px;
    overflow: hidden;
    border: 1px solid #3b4252;
    background: #000;
    box-sizing: border-box;
}

.thumbnail-img {
    width: 100%;
    height: 100%;
    object-fit: cover;
    display: block;
}

.asset-card.missing {
    border: 1px dashed #bf616a;
    display: flex;
    flex-direction: column;
    align-items: center;
    justify-content: center;
}

.missing-placeholder {
    font-size: 9px;
    color: #bf616a;
    margin-top: 2px;
}

.asset-badge {
    position: absolute;
    bottom: 0;
    left: 0;
    right: 0;
    background: rgba(0,0,0,0.85);
    color: #88c0d0;
    font-size: 9px;
    text-align: center;
    line-height: 14px;
}

.cell-flex-wrapper {
    display: flex;
    flex-direction: column;
    gap: 6px;
}

.audio-stack {
    display: flex;
    flex-direction: column;
    gap: 4px;
}

.audio-card {
    display: flex;
    align-items: center;
    gap: 6px;
    background: #232731;
    padding: 4px 8px;
    border-radius: 4px;
}

.audio-card.missing {
    opacity: 0.6;
    border: 1px dashed #bf616a;
}

.audio-player {
    height: 30px;
    width: 100%;
    min-width: 180px;
    flex: 1;
}

.asset-badge-audio {
    font-size: 10px;
    color: #a3be8c;
    font-weight: bold;
    min-width: 36px;
    flex-shrink: 0;
}

/* 视频单元格 */
.video-cell-wrapper {
    display: flex;
    flex-direction: column;
    gap: 6px;
}

.video-wrap {
    position: relative;
    width: 100%;
    cursor: pointer;
}

.video-player {
    width: 100%;
    max-height: 110px;
    border-radius: 4px;
    background: #000;
    display: block;
}

.video-expand-btn {
    position: absolute;
    bottom: 4px;
    right: 4px;
    width: 24px;
    height: 24px;
    border: none;
    border-radius: 4px;
    background: rgba(0,0,0,0.7);
    color: #fff;
    font-size: 14px;
    line-height: 24px;
    text-align: center;
    cursor: pointer;
    padding: 0;
    opacity: 0;
    transition: opacity 0.2s;
    z-index: 2;
}

.video-wrap:hover .video-expand-btn {
    opacity: 1;
}

.video-expand-btn:hover {
    background: rgba(0,0,0,0.9);
}

.raw-text-edit {
    background: #16181f;
    border: 1px solid #3e4452;
    color: #a3be8c;
    padding: 4px 6px;
    font-size: 11px;
    font-family: monospace;
    border-radius: 4px;
    width: 100%;
    box-sizing: border-box;
    outline: none;
}

.raw-text-edit:focus {
    border-color: #88c0d0;
    background: #0e1014;
}

.cell-center {
    text-align: center;
}

/* A.镜头号字体缩小一档 */
.shot-input {
    height: calc(95px * var(--row-scale));
    font-size: 14px !important;
    color: #f4f6fa !important;
    font-weight: 500;
}

.area-title {
    height: calc(120px * var(--row-scale));
    resize: vertical;
    font-size: 14px !important;
    color: #f4f6fa !important;
    line-height: 1.5;
}

.dur-input {
    margin-top: 8px;
}

/* F列提示词框：固定高度，内部滚动 */
.prompt-box {
    background: #14161b;
    border: 1px solid #2e3440;
    padding: 8px;
    border-radius: 4px;
    min-height: calc(100px * var(--row-scale));
    max-height: calc(138px * var(--row-scale));
    overflow-y: auto;
    overflow-x: hidden;
    font-size: 13px;
    line-height: 1.6;
    color: #eceff4;
    white-space: pre-wrap !important;
    word-wrap: break-word !important;
    overflow-wrap: break-word !important;
    word-break: break-word;
    width: 100%;
    box-sizing: border-box;
}

/* 图片灯箱 */
#lightboxModal {
    display: none;
    position: fixed;
    z-index: 99999;
    left: 0;
    top: 0;
    width: 100%;
    height: 100%;
    background: rgba(0,0,0,0.88);
    justify-content: center;
    align-items: center;
}

#lightboxModal img {
    max-width: 92%;
    max-height: 92%;
    border-radius: 6px;
    box-shadow: 0 0 20px rgba(0,0,0,0.9);
}

/* 视频灯箱 */
#videoLightboxModal {
    display: none;
    position: fixed;
    z-index: 99999;
    left: 0;
    top: 0;
    width: 100%;
    height: 100%;
    background: rgba(0,0,0,0.88);
    justify-content: center;
    align-items: center;
}

#videoLightboxModal video {
    max-width: 92%;
    max-height: 92%;
    border-radius: 6px;
    box-shadow: 0 0 20px rgba(0,0,0,0.9);
}
"""

# ==============================
# 4. JS 脚本
# ==============================
CUSTOM_JS = """
let wScale = parseFloat(localStorage.getItem('userColScale')) || 1.0;
let hScale = parseFloat(localStorage.getItem('userRowScale')) || 1.0;

function applyScales() {
    const root = document.documentElement;
    root.style.setProperty('--col-scale', wScale);
    root.style.setProperty('--row-scale', hScale);
}

window.adjustWidth = function(delta) {
    wScale = Math.max(0.6, Math.min(2.5, wScale + delta));
    applyScales();
};

window.adjustHeight = function(delta) {
    hScale = Math.max(0.6, Math.min(2.5, hScale + delta));
    applyScales();
};

window.saveLayoutScale = function() {
    localStorage.setItem('userColScale', wScale);
    localStorage.setItem('userRowScale', hScale);
    showToast('当前栏宽与行高已保存为默认！', 'success');
};

window.resetLayoutScale = function() {
    wScale = 1.0;
    hScale = 1.0;
    localStorage.removeItem('userColScale');
    localStorage.removeItem('userRowScale');
    applyScales();
    showToast('已恢复默认排版', 'success');
};

window.fitWidth = function() {
    const viewport = document.getElementById('storyboardViewport');
    if (!viewport) return;
    const containerWidth = viewport.clientWidth - 20;
    const table = viewport.querySelector('table');
    if (!table) return;
    const tableWidth = table.scrollWidth;
    wScale = containerWidth / tableWidth * wScale;
    wScale = Math.max(0.6, Math.min(1.5, wScale));
    applyScales();
};

window.showToast = function(msg, type = 'success') {
    let toast = document.getElementById('globalToast');
    if (!toast) {
        toast = document.createElement('div');
        toast.id = 'globalToast';
        toast.className = 'toast-box';
        document.body.appendChild(toast);
    }
    toast.className = 'toast-box toast-' + type;
    toast.textContent = msg;
    setTimeout(() => toast.classList.add('toast-show'), 10);
    clearTimeout(window._toastTimer);
    window._toastTimer = setTimeout(() => {
        toast.classList.remove('toast-show');
    }, 2500);
};

window.onRowClick = function(tr) {
    document.querySelectorAll('tr.storyboard-row').forEach(r => r.classList.remove('selected-row'));
    tr.classList.add('selected-row');
    const promptDiv = tr.querySelector('[data-col="prompt"]');
    const box = document.getElementById('masterPromptBox');
    if (promptDiv && box) {
        const text = promptDiv.innerText || '';
        box.innerHTML = "<b style='color:#ebcb8b'>[本镜提示词]：</b>" + (text ? text : "<i style='color:#8c9ba5'>无提示词内容</i>");
    }
};

window.openLightbox = function(url) {
    let modal = document.getElementById('lightboxModal');
    if (!modal) {
        modal = document.createElement('div');
        modal.id = 'lightboxModal';
        modal.innerHTML = '<img id="lightboxImg" src="">';
        modal.onclick = function() { modal.style.display = 'none'; };
        document.body.appendChild(modal);
    }
    document.getElementById('lightboxImg').src = url;
    modal.style.display = 'flex';
};

window.openVideoLightbox = function(src, btnEl) {
    event.stopPropagation();
    const wrap = btnEl.closest('.video-wrap');
    const smallVideo = wrap ? wrap.querySelector('.video-player') : null;

    let currentTime = 0;
    if (smallVideo) {
        currentTime = smallVideo.currentTime;
        smallVideo.pause();
    }

    let modal = document.getElementById('videoLightboxModal');
    if (!modal) {
        modal = document.createElement('div');
        modal.id = 'videoLightboxModal';
        modal.innerHTML = '<video id="lightboxVideo" controls preload="auto" onclick="event.stopPropagation()"></video>';
        modal.onclick = function() {
            const v = document.getElementById('lightboxVideo');
            if (v) v.pause();
            modal.style.display = 'none';
        };
        document.body.appendChild(modal);
    }

    const video = document.getElementById('lightboxVideo');
    video.src = src;
    video.currentTime = currentTime;
    modal.style.display = 'flex';
    video.play().catch(() => {});
};

// 修复：勾选显示，不勾选隐藏
window.applyColVisibility = function(selectedList) {
    const colMap = {
        "镜头号": "col-shot",
        "小标题": "col-title",
        "时长": "col-dur",
        "图片资产": "col-img",
        "声音资产": "col-audio",
        "提示词": "col-prompt",
        "视频结果": "col-video"
    };
    // 先全部隐藏
    Object.values(colMap).forEach(cls => {
        document.querySelectorAll('.' + cls).forEach(el => el.style.display = 'none');
    });
    // 选中的显示
    selectedList.forEach(name => {
        const cls = colMap[name];
        if (cls) {
            document.querySelectorAll('.' + cls).forEach(el => el.style.display = '');
        }
    });
};

window.saveToExcel = function() {
    const viewport = document.getElementById('storyboardViewport');
    const sheetName = viewport ? viewport.dataset.sheetName : '';

    if (!sheetName) {
        showToast('请先加载分镜表', 'error');
        return;
    }

    let rows = [];
    document.querySelectorAll('#storyboardTable tbody tr').forEach(tr => {
        rows.push({
            shot: tr.querySelector('[data-col="shot"]')?.value || "",
            title: tr.querySelector('[data-col="title"]')?.value || "",
            dur: tr.querySelector('[data-col="dur"]')?.value || "0",
            img: tr.querySelector('[data-col="img"]')?.value || "",
            audio: tr.querySelector('[data-col="audio"]')?.value || "",
            prompt: tr.querySelector('[data-col="prompt"]')?.innerText || "",
            video: tr.querySelector('[data-col="video"]')?.value || ""
        });
    });

    if (rows.length === 0) {
        showToast('表格数据为空', 'error');
        return;
    }

    fetch('/api/save_excel', {
        method: 'POST',
        headers: {'Content-Type': 'application/json'},
        body: JSON.stringify({
            sheet_name: sheetName,
            rows: rows
        })
    }).then(res => res.json()).then(res => {
        if (res.status === 'success') {
            showToast(res.msg, 'success');
        } else {
            showToast(res.msg, 'error');
        }
    }).catch(err => {
        showToast('保存失败：网络错误', 'error');
    });
};

window.pauseOtherVideos = function(currentVideo) {
    document.querySelectorAll('.video-player').forEach(v => {
        if (v !== currentVideo) v.pause();
    });
};

document.addEventListener('keydown', function(e) {
    const activeTag = document.activeElement.tagName;
    if (activeTag === 'INPUT' || activeTag === 'TEXTAREA' || activeTag === 'SELECT') return;
    const step = e.shiftKey ? 0.02 : 0.1;

    if (e.key === 'g' || e.key === 'G') {
        e.preventDefault();
        adjustWidth(-step);
    }
    if (e.key === 'h' || e.key === 'H') {
        e.preventDefault();
        adjustWidth(step);
    }
    if (e.key === 'y' || e.key === 'Y') {
        e.preventDefault();
        adjustHeight(-step);
    }
    if (e.key === 'u' || e.key === 'U') {
        e.preventDefault();
        adjustHeight(step);
    }
    if (e.key === 'r' || e.key === 'R') {
        e.preventDefault();
        resetLayoutScale();
    }
    if (e.key === 'a' || e.key === 'A') {
        e.preventDefault();
        const btn = Array.from(document.querySelectorAll('button')).find(b => b.textContent.includes('上一场'));
        if (btn) btn.click();
    }
    if (e.key === 's' || e.key === 'S') {
        e.preventDefault();
        const btn = Array.from(document.querySelectorAll('button')).find(b => b.textContent.includes('下一场'));
        if (btn) btn.click();
    }
});

document.addEventListener('wheel', function(e) {
    const viewport = document.getElementById('storyboardViewport');
    if (!viewport) return;
    if (!viewport.contains(e.target)) return;
    if (e.shiftKey) {
        e.preventDefault();
        viewport.scrollLeft += e.deltaY;
    }
}, { passive: false });

setTimeout(applyScales, 500);
"""

# ==============================
# 5. 后端控制器
# ==============================
class StoryboardApp:
    def __init__(self):
        self.excel_b_path, self.asset_dir, self.video_dir = "", "", ""
        self.image_index = {}
        self.audio_index = {}
        self.video_paths = []
        self.b_sheets, self.current_sheet_b = [], ""

    def load_project(self, path_a, path_b, asset_p, video_p, current_sheet_val):
        save_config(path_a, path_b, asset_p, video_p)
        self.excel_b_path = clean_path(path_b)
        self.asset_dir = clean_path(asset_p)
        self.video_dir = clean_path(video_p)
        self.image_index = build_image_index(self.asset_dir)
        self.audio_index = build_audio_index(self.asset_dir)
        target_video_dir = self.video_dir if self.video_dir else self.asset_dir
        self.video_paths = build_video_list(target_video_dir)

        sheet_choices, default_sheet = [], None
        if os.path.exists(self.excel_b_path):
            try:
                wb = openpyxl.load_workbook(self.excel_b_path, read_only=True)
                self.b_sheets = sorted(wb.sheetnames, key=natural_sort_key)
                sheet_choices = self.b_sheets
                default_sheet = current_sheet_val if current_sheet_val in sheet_choices else (self.b_sheets[0] if self.b_sheets else None)
                self.current_sheet_b = default_sheet
            except Exception as e:
                gr.Warning(f"读取分镜表失败: {e}")
        return gr.update(choices=sheet_choices, value=default_sheet)

    def load_sheet_b_content(self, sheet_name):
        if not self.excel_b_path or not os.path.exists(self.excel_b_path) or not sheet_name:
            return "<div class='empty-tip'>请先加载分镜表</div>"
        self.current_sheet_b = sheet_name
        df = pd.read_excel(self.excel_b_path, sheet_name=sheet_name)
        return render_table_b_html(df, self.image_index, self.audio_index, self.video_paths, sheet_name)

    def step_sheet(self, direction, current_sheet):
        if not self.b_sheets or current_sheet not in self.b_sheets:
            return current_sheet, self.load_sheet_b_content(current_sheet)
        new_idx = self.b_sheets.index(current_sheet) + direction
        if 0 <= new_idx < len(self.b_sheets):
            ts = self.b_sheets[new_idx]
            return ts, self.load_sheet_b_content(ts)
        return current_sheet, self.load_sheet_b_content(current_sheet)

app_core = StoryboardApp()
init_conf = load_config()

# ==============================
# 6. UI 页面构建
# ==============================
with gr.Blocks(title="AI 影视分镜助手") as demo:
    with gr.Row():
        txt_path_a = gr.Textbox(label="表格 A (资产表)", value=init_conf.get("path_a", ""), lines=1, max_lines=1, scale=2)
        txt_path_b = gr.Textbox(label="表格 B (分镜表)", value=init_conf.get("path_b", ""), lines=1, max_lines=1, scale=2)
        txt_asset_dir = gr.Textbox(label="资产目录", value=init_conf.get("asset_dir", ""), lines=1, max_lines=1, scale=2)
        txt_video_dir = gr.Textbox(label="视频目录", value=init_conf.get("video_dir", ""), lines=1, max_lines=1, scale=2, placeholder="不填默认同资产目录")
        btn_load_project = gr.Button("🚀 加载/刷新工程", variant="primary", scale=1)

    with gr.Tabs():
        with gr.TabItem("📋 表格 B · 影视分镜表"):

            # 第一行：场次 + 隐藏列 + 上下场按钮
            with gr.Row(elem_classes="top-toolbar-row"):
                dd_sheets_b = gr.Dropdown(label="当前场次", choices=[], interactive=True, scale=2)
                with gr.Column(scale=3, elem_classes="col-checkbox-wrap"):
                    chk_cols = gr.CheckboxGroup(
                        choices=["镜头号", "小标题", "时长", "图片资产", "声音资产", "提示词", "视频结果"],
                        value=["镜头号", "小标题", "时长", "图片资产", "声音资产", "提示词", "视频结果"],
                        label="✓ 隐藏/显示列",
                        scale=0
                    )
                btn_prev_sheet = gr.Button("◀ 上一场", scale=0, min_width=60)
                btn_next_sheet = gr.Button("下一场 ▶", scale=0, min_width=60)

            # 第二行：视图微调 横向平铺
            with gr.Row(elem_classes="view-toolbar-row"):
                gr.Markdown("**📐 视图微调**", scale=0)
                btn_w_dec = gr.Button("◀ 栏宽收窄", scale=1)
                btn_w_inc = gr.Button("▶ 栏宽拉宽", scale=1)
                btn_h_dec = gr.Button("▲ 行高收紧", scale=1)
                btn_h_inc = gr.Button("▼ 行高拉开", scale=1)
                btn_fit_width = gr.Button("↔ 适配宽度", scale=1)
                btn_scale_save = gr.Button("📌 记住排版", scale=1)
                btn_scale_reset = gr.Button("🔄 恢复默认", scale=1)

            # 第三行：主提示词 + 保存按钮
            with gr.Row(elem_classes="prompt-save-row"):
                gr.HTML("""<div id="masterPromptBox" class="master-prompt-preview"></div>""")
                btn_save_to_excel = gr.Button("💾 同步并保存到 Excel", variant="primary", scale=0, elem_classes="save-btn")

            html_storyboard_view = gr.HTML("<div class='empty-tip'>等待加载工程数据...</div>")

    # ==============================
    # 事件绑定
    # ==============================
    btn_w_dec.click(None, None, None, js="adjustWidth(-0.1)")
    btn_w_inc.click(None, None, None, js="adjustWidth(0.1)")
    btn_h_dec.click(None, None, None, js="adjustHeight(-0.1)")
    btn_h_inc.click(None, None, None, js="adjustHeight(0.1)")
    btn_fit_width.click(None, None, None, js="fitWidth()")
    btn_scale_save.click(None, None, None, js="saveLayoutScale()")
    btn_scale_reset.click(None, None, None, js="resetLayoutScale()")

    chk_cols.change(None, inputs=[chk_cols], js="(cols) => { applyColVisibility(cols); }")

    btn_load_project.click(fn=app_core.load_project, inputs=[txt_path_a, txt_path_b, txt_asset_dir, txt_video_dir, dd_sheets_b], outputs=[dd_sheets_b])
    dd_sheets_b.change(fn=app_core.load_sheet_b_content, inputs=[dd_sheets_b], outputs=[html_storyboard_view])
    btn_prev_sheet.click(fn=lambda cur: app_core.step_sheet(-1, cur), inputs=[dd_sheets_b], outputs=[dd_sheets_b, html_storyboard_view])
    btn_next_sheet.click(fn=lambda cur: app_core.step_sheet(1, cur), inputs=[dd_sheets_b], outputs=[dd_sheets_b, html_storyboard_view])

    btn_save_to_excel.click(None, None, None, js="saveToExcel()")

# ==============================
# 7. FastAPI 接口
# ==============================
fastapi_app = FastAPI()

class SaveExcelRequest(BaseModel):
    sheet_name: str
    rows: list

@fastapi_app.post("/api/save_excel")
def save_excel_api(req: SaveExcelRequest):
    try:
        sheet_name = req.sheet_name
        rows = req.rows

        if not app_core.excel_b_path or not os.path.exists(app_core.excel_b_path):
            return {"status": "error", "msg": "保存失败：Excel 文件不存在"}
        if not sheet_name or not rows:
            return {"status": "error", "msg": "保存失败：数据不完整"}

        backup_excel(app_core.excel_b_path)

        wb = openpyxl.load_workbook(app_core.excel_b_path)
        ws = wb[sheet_name]
        for idx, row_data in enumerate(rows):
            r = idx + 2
            ws.cell(row=r, column=1, value=row_data.get('shot', ''))
            ws.cell(row=r, column=2, value=row_data.get('title', ''))
            try: ws.cell(row=r, column=3, value=float(row_data.get('dur', 0)))
            except: pass
            ws.cell(row=r, column=4, value=row_data.get('img', ''))
            ws.cell(row=r, column=5, value=row_data.get('audio', ''))
            ws.cell(row=r, column=6, value=row_data.get('prompt', ''))
            ws.cell(row=r, column=7, value=row_data.get('video', ''))
        wb.save(app_core.excel_b_path)
        filename = os.path.basename(app_core.excel_b_path)
        return {"status": "success", "msg": f"✅ 已保存到 {filename}"}

    except PermissionError:
        return {"status": "error", "msg": "❌ 保存失败：Excel 文件正在被占用，请关闭 Excel/WPS 后重试"}
    except Exception as e:
        return {"status": "error", "msg": f"❌ 保存失败: {str(e)}"}

@fastapi_app.get("/api/local_media")
def serve_local_media(filepath: str):
    path = urllib.parse.unquote(filepath)
    if not os.path.exists(path):
        return Response(content=b"File not found", status_code=404)

    ext = os.path.splitext(path)[1].lower()

    if ext in {'.jpg', '.jpeg', '.png', '.webp', '.bmp', '.tif', '.tiff'}:
        try:
            with Image.open(path) as img:
                try:
                    img = ImageOps.exif_transpose(img)
                except Exception:
                    pass
                if img.mode not in ('RGB', 'L'):
                    img = img.convert('RGB')
                buf = io.BytesIO()
                img.save(buf, format="JPEG", quality=90, optimize=True)
                buf.seek(0)
                return Response(content=buf.getvalue(), media_type="image/jpeg")
        except Exception:
            pass

    mime_type, _ = mimetypes.guess_type(path)
    return FileResponse(path, media_type=mime_type or "application/octet-stream")

app = gr.mount_gradio_app(fastapi_app, demo, path="/", css=CUSTOM_CSS, js=CUSTOM_JS)

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=7861)
