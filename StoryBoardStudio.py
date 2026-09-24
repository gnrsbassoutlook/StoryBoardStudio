import os
import re
import io
import sys
import json
import html
import shutil
import asyncio
import mimetypes
import subprocess
from datetime import datetime
import urllib.parse
import pandas as pd
import openpyxl
from PIL import Image, ImageOps, ImageFile
import gradio as gr
from fastapi import FastAPI, Response
from fastapi.responses import FileResponse
from pydantic import BaseModel

ImageFile.LOAD_TRUNCATED_IMAGES = True

# 第一部分（资产表预览）—— 独立模块，不改动下方第二部分任何逻辑
from asset_tab import ASSET_CSS, ASSET_JS, build_asset_tab, register_asset_api

# ==============================
# 0. 本地配置持久化
# ==============================
CONFIG_DIR = os.path.join(os.getcwd(), "gradio_temp")
CONFIG_FILE = os.path.join(CONFIG_DIR, "config.json")

def load_config():
    if os.path.exists(CONFIG_FILE):
        try:
            with open(CONFIG_FILE, 'r', encoding='utf-8') as f:
                return json.load(f)
        except:
            pass
    return {"path_a": "", "path_b": "", "asset_dir": "", "video_dir": "", "audio_dir": ""}

def save_config(path_a, path_b, asset_dir, video_dir, audio_dir):
    os.makedirs(CONFIG_DIR, exist_ok=True)
    with open(CONFIG_FILE, 'w', encoding='utf-8') as f:
        json.dump({
            "path_a": path_a,
            "path_b": path_b,
            "asset_dir": asset_dir,
            "video_dir": video_dir,
            "audio_dir": audio_dir
        }, f, ensure_ascii=False, indent=2)

# ==============================
# 1. 资产索引与冲突检测（精准匹配X/Y/Z编码）
# ==============================
# 「复制文件地址」可能带上的引号：资源管理器给的是英文双引号，某些输入法 /
# 聊天工具 / 文档里粘出来的是中文引号，这里一律剥掉。
_QUOTE_CHARS = "\"'“”‘’「」『』"

def _trim_sep(p: str) -> str:
    """去掉路径尾部的分隔符（macOS 选文件夹会带 /），但保住盘符根目录与 / 。"""
    if not p or p == "/" or re.match(r"^[A-Za-z]:[\\/]?$", p):
        return p
    return p.rstrip("/\\")

def clean_path(path_str: str) -> str:
    """把用户粘进来的路径洗成干净的绝对路径。

    兼容这些「模糊输入」：
      · 首尾空格、全角空格、BOM
      · 被引号包裹 —— 资源管理器「复制文件地址」会给 "C:\\a\\b.xlsx" 套一对英文
        双引号，中文引号 “ ” 和单引号 ' ' 同理；只粘到左引号（右引号被吃掉）也能处理
      · 浏览器里复制来的 file:// 前缀
      · 终端里拖文件产生的「反斜杠 + 空格」转义
    """
    if not path_str:
        return ""
    p = str(path_str).lstrip("\ufeff").replace("\u3000", " ").strip()
    if p[:7].lower() == "file://":
        p = urllib.parse.unquote(p[7:])
        if re.match(r"^/[A-Za-z]:", p):        # file:///C:/x → C:/x
            p = p[1:]
    for _ in range(3):                          # 反复剥，兼容 外层"内层' 这种嵌套
        p = p.strip().strip("\u200b")
        if len(p) >= 2 and p[0] in _QUOTE_CHARS and p[-1] in _QUOTE_CHARS:
            p = p[1:-1]
        elif p and p[0] in _QUOTE_CHARS:        # 只粘到左引号的情况
            p = p[1:]
        else:
            break
    p = p.strip().replace("\\ ", " ")
    if not p:
        return ""
    return os.path.abspath(p)

def natural_sort_key(s: str):
    return [int(text) if text.isdigit() else text.lower() for text in re.split(r'(\d+)', str(s))]

def _build_index_and_conflicts(root_dir: str, valid_exts: set):
    index = {}
    conflicts = []
    if not root_dir or not os.path.isdir(root_dir):
        return index, conflicts
    
    code_groups = {}
    # 仅匹配X/Y/Z开头+数字的资产编码，可选后缀字母
    prefix_pattern = re.compile(r"^([XYZxyz]\d+[a-zA-Z]?)[._\-]")
    
    for root, _, files in os.walk(root_dir):
        for f in files:
            ext = os.path.splitext(f)[1].lower()
            if ext not in valid_exts:
                continue
            
            full_path = os.path.join(root, f)
            file_name_no_ext = os.path.splitext(f)[0]
            
            # 前缀编码匹配（仅X/Y/Z开头的标准资产编码）
            match = prefix_pattern.match(f)
            if match:
                prefix_code = match.group(1).upper()
                if prefix_code not in code_groups:
                    code_groups[prefix_code] = []
                code_groups[prefix_code].append((f, full_path))
            
            # 完整文件名匹配（支持镜头号命名，不参与冲突检测）
            full_code = file_name_no_ext.upper()
            if full_code not in index:
                index[full_code] = full_path
    
    # 生成索引与冲突列表
    for code, file_list in code_groups.items():
        file_list.sort(key=lambda x: natural_sort_key(x[0]))
        index[code] = file_list[0][1]
        if len(file_list) > 1:
            conflict_names = "、".join([f[0] for f in file_list])
            conflicts.append(f"{code} → {conflict_names}")
    
    return index, conflicts

def build_image_index(asset_dir: str):
    valid_exts = {'.jpg', '.jpeg', '.png', '.webp', '.bmp', '.tif', '.tiff'}
    return _build_index_and_conflicts(asset_dir, valid_exts)

def build_audio_index(audio_dir: str):
    valid_exts = {'.mp3', '.wav', '.aac', '.m4a', '.flac'}
    return _build_index_and_conflicts(audio_dir, valid_exts)

def build_video_list(video_dir: str):
    video_paths = []
    if not video_dir or not os.path.isdir(video_dir):
        return video_paths
    video_exts = {'.mp4', '.mov', '.webm', '.mkv'}
    for root, _, files in os.walk(video_dir):
        for f in files:
            ext = os.path.splitext(f)[1].lower()
            if ext in video_exts:
                video_paths.append(os.path.join(root, f))
    return video_paths

def find_matching_video(shot_id: str, video_paths: list):
    shot_lower = shot_id.lower()
    separators = {'.', '-', '_'}
    for path in video_paths:
        fname = os.path.basename(path)
        name_body, _ = os.path.splitext(fname)
        name_lower = name_body.lower()
        if name_lower == shot_lower:
            return path
        if any(name_lower.startswith(f"{shot_lower}{sep}") for sep in separators):
            return path
    return None

# C列「严格代表时长」（模式2）的写法：12 / 12.5 / 12f / 12s（容忍数字与 f/s 之间多余空格）
DURATION_PATTERN = re.compile(r'^\d+(\.\d+)?\s*[fFsS]?$')
# 镜头号与描述之间的分隔符（口径与 find_matching_video 一致）
_MATCH_SEPARATORS = {'.', '-', '_', ' '}

def is_duration_value(text: str) -> bool:
    """C列是否「严格代表时长」（模式2）。

    只有纯数字、数字+f、数字+s 才算（数字可以是任意浮点数：12、2.8、12.75、0.5）。
    `n/a`、`mode_1`、`mode_A`、`模式1` 这类模式1 的标记不是数字，不会被这里吃掉 ——
    它们会原样显示在单元格里；也不会有任何值让界面报错。
    """
    return bool(DURATION_PATTERN.match((text or "").strip()))

def format_duration_text(text: str) -> str:
    """时长的显示文本：`8.0` 这种「整数被 Excel 存成浮点」的情况去掉尾巴，其余原样。

    只影响界面显示，**不改动 Excel 里的值**（写回时用的仍是单元格原始内容）。
    `2.8` 这类真小数不受影响。
    """
    s = (text or "").strip()
    m = re.match(r'^(\d+)\.0+(\s*[fFsS]?)$', s)
    return (m.group(1) + m.group(2)) if m else s

def find_matching_audio(shot_id: str, audio_index: dict):
    """按镜头号在音频索引里找音频文件，返回绝对路径；找不到返回空串。

    `audio_index` 的键是「去掉扩展名的文件名 .upper()」（mp3/wav/flac/aac/m4a 都收录）。
    两级匹配：
      1. 精确：`s01-01.mp3` → 键 `S01-01`
      2. 前缀模糊：`s01-01.何功伟的话.mp3` → 键以 `S01-01` + 分隔符 开头
    多个模糊命中时按自然序取第一个，保证每次渲染结果稳定。
    """
    key = (shot_id or "").strip().upper()
    if not key or key == "0" or not audio_index:
        return ""
    if key in audio_index:
        return audio_index[key]
    hits = [k for k in audio_index
            if k.startswith(key) and len(k) > len(key) and k[len(key)] in _MATCH_SEPARATORS]
    if hits:
        hits.sort(key=natural_sort_key)
        return audio_index[hits[0]]
    return ""

def backup_excel(filepath: str, asset_dir: str = ""):
    """保存写回前先备份原 Excel。

    备份位置：**资产目录/excel_backup/**（用户口径：Excel 与它的历史版本归一堆，好找）。
    资产目录没填 / 不存在时，退回老位置 —— Excel 同目录的 `_backups/`。
    返回备份文件路径；文件不存在时返回空串。
    """
    if not os.path.exists(filepath):
        return ""
    base = clean_path(asset_dir) if asset_dir else ""
    if base and os.path.isdir(base):
        backup_dir = os.path.join(base, "excel_backup")
    else:
        backup_dir = os.path.join(os.path.dirname(filepath), "_backups")
    os.makedirs(backup_dir, exist_ok=True)
    basename = os.path.basename(filepath)
    name, ext = os.path.splitext(basename)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = os.path.join(backup_dir, f"{name}_{timestamp}{ext}")
    shutil.copy2(filepath, backup_path)
    return backup_path

# ==============================
# 2. 核心表格渲染引擎（D/E列互换）
# ==============================
def render_table_b_html(df, image_index, audio_index, video_paths, sheet_name):
    if df is None or df.empty:
        return "<div class='empty-tip'>暂无数据或工作表为空</div>"

    rows_html = []

    for idx, row in df.iterrows():
        shot_id = str(row.iloc[0]) if len(row) > 0 and pd.notna(row.iloc[0]) else f"Sht{idx+1:02d}"
        title = str(row.iloc[1]) if len(row) > 1 and pd.notna(row.iloc[1]) else ""
        
        # C列：双模式兼容
        #  · 纯数字 / 数字+f / 数字+s → 严格视为「时长」（模式2），只显示数字本身，不找音频
        #  · n/a、mode_1、模式1 等     → 模式1 的标记，原样显示，不当时长解析
        #  · 其余（通常是镜头号）       → 模式1，原样显示，并按镜头号模糊匹配音频挂播放条
        c_content = str(row.iloc[2]) if len(row) > 2 and pd.notna(row.iloc[2]) else ""
        c_content_strip = c_content.strip()
        
        c_cell_top = ""
        if is_duration_value(c_content_strip):
            c_cell_top = f"""
            <div class="dur-text">
                <span class="dur-value">{html.escape(format_duration_text(c_content_strip))}</span>
            </div>"""
        else:
            # 模式1：匹配候选 C 列自身优先（常直接填镜头号），A 列镜头号兜底（填 n/a 等标记时）
            audio_path = ""
            for candidate in (c_content_strip, shot_id):
                audio_path = find_matching_audio(candidate, audio_index)
                if audio_path:
                    break
            if audio_path:
                file_url = f"/api/local_media?filepath={urllib.parse.quote(audio_path)}"
                badge_text = html.escape(c_content_strip or shot_id)
                c_cell_top = f"""
                <div class="audio-card single-audio">
                    <span class="asset-badge-audio">{badge_text}</span>
                    <audio controls preload="auto" src="{file_url}" class="audio-player" onclick="event.stopPropagation()" playsinline>
                </div>"""
            else:
                c_cell_top = f"""
                <div class="dur-text">
                    <span class="dur-value">{html.escape(c_content_strip)}</span>
                </div>"""
        
        c_cell = f"""
        <div class="cell-flex-wrapper">
            {c_cell_top}
            <input type="text" class="raw-text-edit" data-col="dur" value="{html.escape(c_content)}" placeholder="音频编码/时长"/>
        </div>"""

        # D列：台词内容/角色
        dialogue_text = str(row.iloc[3]) if len(row) > 3 and pd.notna(row.iloc[3]) else ""
        dialogue_cell = f"""
        <div class="cell-flex-wrapper">
            <textarea class="dialogue-input" data-col="dialogue" placeholder="台词内容/角色">{dialogue_text}</textarea>
        </div>"""

        # E列：图片资产
        img_codes_str = str(row.iloc[4]) if len(row) > 4 and pd.notna(row.iloc[4]) else ""
        img_items = [c.strip() for c in img_codes_str.split('|') if c.strip() and c.strip() != '0']
        img_previews = []
        for code in img_items[:6]:
            code_up = code.upper()
            if code_up in image_index:
                file_url = f"/api/local_media?filepath={urllib.parse.quote(image_index[code_up])}"
                img_previews.append(f"""
                    <div class="asset-card" onclick="openLightbox('{file_url}')" title="点击放大: {code}">
                        <img src="{file_url}" alt="{code}" class="thumbnail-img"/>
                        <span class="asset-badge">{code}</span>
                    </div>""")
            else:
                img_previews.append(f"""<div class="asset-card missing" title="未找到文件"><span class="asset-badge">{code}</span><div class="missing-placeholder">无</div></div>""")

        img_grid = f"<div class='img-grid'>{''.join(img_previews)}</div>"
        img_cell_content = f"""
        <div class="cell-flex-wrapper">
            {img_grid}
            <input type="text" class="raw-text-edit img-code-input" value="{img_codes_str}" data-col="img" placeholder="图片编码"/>
        </div>"""

        prompt = str(row.iloc[5]) if len(row) > 5 and pd.notna(row.iloc[5]) else ""

        vid_path = find_matching_video(shot_id, video_paths)

        # G列视频
        video_preview = ""
        if vid_path:
            vid_url = f"/api/local_media?filepath={urllib.parse.quote(vid_path)}"
            video_preview = f"""
            <div class="cell-flex-wrapper">
                <div class="video-wrap" 
                     onmouseenter="this.querySelector('video').controls = true"
                     onmouseleave="this.querySelector('video').controls = false">
                    <video class="video-player"
                           preload="metadata"
                           playsinline
                           onplay="pauseOtherVideos(this)">
                        <source src="{vid_url}">
                    </video>
                    <button class="video-expand-btn" onclick="openVideoLightbox('{vid_url}', this)" title="放大查看">⛶</button>
                </div>
                <div class="video-label">自动匹配</div>
            </div>
            """
        else:
            video_preview = f"""
            <div class="cell-flex-wrapper">
                <div class="video-empty">
                    <span>自动匹配</span>
                </div>
            </div>
            """

        rows_html.append(f"""
            <tr data-row-idx="{idx}" class="storyboard-row" onclick="onRowClick(this)">
                <td class="col-shot">
                    <input type="text" class="raw-text-edit cell-center shot-input" data-col="shot" value="{shot_id}"/>
                </td>
                <td class="col-title">
                    <textarea class="raw-text-edit area-title" data-col="title">{title}</textarea>
                </td>
                <td class="col-dur">{c_cell}</td>
                <td class="col-dialogue">{dialogue_cell}</td>
                <td class="col-img">{img_cell_content}</td>
                <td class="col-prompt">
                    <textarea class="prompt-box" data-col="prompt">{prompt}</textarea>
                </td>
                <td class="col-video">{video_preview}</td>
            </tr>
        """)

    table_body = "".join(rows_html)

    return f"""
    <div class="storyboard-viewport" id="storyboardViewport" data-sheet-name="{sheet_name}">
        <table class="storyboard-table" id="storyboardTable">
            <thead>
                <tr>
                    <th class="col-shot">A.镜头号</th>
                    <th class="col-title">B.小标题</th>
                    <th class="col-dur">C.音频结果（模式1）/镜头时长（模式2）</th>
                    <th class="col-dialogue">D.台词内容（模式1）/台词角色（模式2）</th>
                    <th class="col-img">E.图片资产</th>
                    <th class="col-prompt">F.生视频提示词</th>
                    <th class="col-video">G.视频结果</th>
                </tr>
            </thead>
            <tbody>{table_body}</tbody>
        </table>
    </div>
    """

# ==============================
# 3. CSS 样式
# ==============================
CUSTOM_CSS = """
:root {
    --col-scale: 1.0;
    --row-scale: 1.0;
    --img-base-w: 82px;
    --img-base-h: 58px;
    --header-height: 36px;
}

.hidden-element { display: none !important; }

.gradio-container {
    max-width: 98% !important;
    background-color: #121418 !important;
    color: #e1e4ea !important;
}

/* 冲突提示区 */
.conflict-panel {
    padding: 4px 10px;
    background: #2a1618;
    border: 1px solid #5c2e2e;
    border-radius: 6px;
    margin-bottom: 6px;
    max-height: 60px;
    overflow-y: auto;
    overflow-x: hidden;
    font-size: 11px;
    color: #f08890;
    line-height: 1.5;
}
.conflict-panel:empty { display: none; }
.conflict-panel .conflict-item {
    display: block;
    margin-bottom: 2px;
}

/* Toast 提示 */
.toast-box {
    position: fixed;
    top: 20px;
    left: 50%;
    transform: translateX(-50%);
    z-index: 10000;
    padding: 10px 20px;
    border-radius: 8px;
    font-size: 13px;
    font-weight: 500;
    box-shadow: 0 4px 12px rgba(0,0,0,0.3);
    opacity: 0;
    transition: opacity 0.3s, top 0.3s;
    pointer-events: none;
}
.toast-success { background: #2d5a3d; color: #d4edda; border: 1px solid #3d7a4f; }
.toast-error { background: #5c2e2e; color: #f8d7da; border: 1px solid #7a3a3a; }
.toast-show { opacity: 1; top: 30px; }

/* 顶部工具栏 */
.top-toolbar-row {
    display: flex;
    align-items: center;
    gap: 8px;
    padding: 6px 10px;
    background: #1c2027;
    border: 1px solid #2b303c;
    border-radius: 8px;
    margin-bottom: 6px;
}

.col-checkbox-wrap {
    flex: 1;
    display: flex;
    align-items: center;
}

.view-toolbar-row {
    display: flex;
    align-items: center;
    gap: 6px;
    padding: 4px 10px;
    background: #1c2027;
    border: 1px solid #2b303c;
    border-radius: 8px;
    margin-bottom: 6px;
    flex-wrap: nowrap;
}

.view-toolbar-row button {
    min-height: 28px !important;
    font-size: 11px !important;
    padding: 0 8px !important;
    flex-shrink: 0;
}

/* 主提示区 + 保存按钮 */
.prompt-save-row {
    display: flex;
    align-items: stretch;
    gap: 8px;
    margin-bottom: 6px;
}

.master-prompt-preview {
    flex: 1;
    padding: 8px 10px;
    background: #1c2027;
    border-left: 3px solid #ebcb8b;
    border-radius: 4px;
    height: 90px;
    box-sizing: border-box;
    overflow-y: auto;
    overflow-x: hidden;
    font-size: 14px;
    line-height: 1.5;
    color: #f0f3f8;
    white-space: pre-wrap !important;
    word-wrap: break-word !important;
}

.save-btn {
    min-height: 28px !important;
    align-self: stretch;
    min-width: 160px !important;
    font-size: 13px !important;
}

/* 表格视口 */
.storyboard-viewport {
    width: 100%;
    overflow: auto;
    max-height: 72vh;
    border: 1px solid #2b303c;
    border-radius: 8px;
    background: #181a20;
    scroll-behavior: smooth;
}

.storyboard-table {
    width: max-content;
    min-width: 100%;
    border-collapse: collapse;
    font-size: 12px;
    table-layout: fixed;
}

.storyboard-table thead th {
    position: sticky;
    top: 0;
    background: #20242c;
    color: #c8d1e0;
    font-weight: 600;
    border-bottom: 2px solid #2e3440;
    z-index: 10;
    text-align: left;
    height: var(--header-height);
    line-height: var(--header-height);
    padding: 0 8px;
    white-space: nowrap;
    font-size: 12px;
}

.storyboard-table td {
    border-bottom: 1px solid #282c37;
    padding: 4px 6px;
    vertical-align: top;
}

.storyboard-row:hover {
    background-color: #1a1e27;
    cursor: pointer;
}

.storyboard-row.selected-row {
    background-color: #242e3f !important;
    box-shadow: inset 4px 0 0 #88c0d0;
}

/* 列宽 */
.col-shot { width: calc(72px * var(--col-scale)); }
.col-title { width: calc(130px * var(--col-scale)); }
.col-dur { width: calc(220px * var(--col-scale)); }
.col-dialogue { width: calc(190px * var(--col-scale)); }
.col-img { width: calc(290px * var(--col-scale)); }
.col-prompt { width: calc(330px * var(--col-scale)); }
.col-video { width: calc(230px * var(--col-scale)); }

/* D列台词 */
.dialogue-input {
    background: #14161b;
    border: 1px solid #2e3440;
    padding: 6px;
    border-radius: 4px;
    font-size: 12px;
    line-height: 1.4;
    color: #c8d1e0;
    resize: vertical;
    width: 100%;
    box-sizing: border-box;
    min-height: calc(104px * var(--row-scale));
    max-height: calc(144px * var(--row-scale));
    font-family: inherit;
}

/* E列图片 */
.img-grid {
    display: grid;
    grid-template-columns: repeat(3, calc(var(--img-base-w) * var(--col-scale)));
    gap: 4px;
    width: 100%;
}

.asset-card {
    position: relative;
    width: calc(var(--img-base-w) * var(--col-scale));
    height: calc(var(--img-base-h) * var(--col-scale));
    border-radius: 3px;
    overflow: hidden;
    border: 1px solid #3b4252;
    background: #0a0a0f;
    box-sizing: border-box;
    flex-shrink: 0;
}

.thumbnail-img {
    width: 100%;
    height: 100%;
    object-fit: contain;
    display: block;
}

.asset-card.missing {
    border: 1px dashed #bf616a;
    display: flex;
    flex-direction: column;
    align-items: center;
    justify-content: center;
}

.missing-placeholder {
    font-size: 9px;
    color: #bf616a;
    margin-top: 2px;
}

.asset-badge {
    position: absolute;
    bottom: 0;
    left: 0;
    right: 0;
    background: rgba(0,0,0,0.85);
    color: #88c0d0;
    font-size: 9px;
    text-align: center;
    line-height: 14px;
}

.img-code-input {
    font-size: 10px !important;
    padding: 2px 6px !important;
    margin-top: 2px;
}

.cell-flex-wrapper {
    display: flex;
    flex-direction: column;
    gap: 3px;
}

/* C列 */
.single-audio {
    margin-top: 2px;
}

.dur-text {
    display: flex;
    align-items: baseline;
    gap: 4px;
    padding: 6px 6px;
    background: #232731;
    border-radius: 4px;
    margin-top: 2px;
}
.dur-value {
    font-size: 15px;
    font-weight: 600;
    color: #ebcb8b;
}

.audio-card {
    display: flex;
    align-items: center;
    gap: 6px;
    background: #232731;
    padding: 3px 6px;
    border-radius: 4px;
}

.audio-player {
    height: 28px;
    width: 100%;
    flex: 1;
    min-width: 120px;
}

.asset-badge-audio {
    font-size: 10px;
    color: #a3be8c;
    font-weight: bold;
    min-width: 36px;
    flex-shrink: 0;
}

/* G列视频 */
.video-cell-wrapper {
    display: flex;
    flex-direction: column;
    gap: 3px;
}

.video-wrap {
    position: relative;
    width: 100%;
    cursor: pointer;
    border-radius: 4px;
    overflow: hidden;
    background: #000;
}

.video-player {
    width: 100%;
    max-height: 100px;
    display: block;
    object-fit: contain;
    background: #000;
}

.video-expand-btn {
    position: absolute;
    bottom: 4px;
    right: 4px;
    width: 22px;
    height: 22px;
    border: none;
    border-radius: 4px;
    background: rgba(0,0,0,0.7);
    color: #fff;
    font-size: 12px;
    line-height: 22px;
    text-align: center;
    cursor: pointer;
    padding: 0;
    opacity: 0;
    transition: opacity 0.2s;
    z-index: 2;
}

.video-wrap:hover .video-expand-btn {
    opacity: 1;
}

.video-label {
    font-size: 10px;
    color: #8c9ba5;
    text-align: center;
}

.video-empty {
    height: calc(100px * var(--row-scale));
    background: #0a0a0f;
    border: 1px solid #2e3440;
    border-radius: 4px;
    display: flex;
    align-items: center;
    justify-content: center;
    color: #6b7280;
    font-size: 11px;
}

/* 通用输入框 */
.raw-text-edit {
    background: #16181f;
    border: 1px solid #3e4452;
    color: #a3be8c;
    padding: 3px 6px;
    font-size: 11px;
    font-family: monospace;
    border-radius: 4px;
    width: 100%;
    box-sizing: border-box;
    outline: none;
}

.raw-text-edit:focus {
    border-color: #88c0d0;
    background: #0e1014;
}

.cell-center {
    text-align: center;
}

.shot-input {
    height: 30px;
    font-size: 13px !important;
    color: #f4f6fa !important;
    font-weight: 500;
    margin-top: 4px;
}

.area-title {
    /* 必须消费 --row-scale：Shift+g / Shift+h 只改这个变量，
       不写进 calc 的话纵向缩放就是空操作 */
    height: calc(106px * var(--row-scale));
    resize: vertical;
    font-size: 13px !important;
    color: #f4f6fa !important;
    line-height: 1.4;
    font-family: inherit;
}

/* F列提示词 */
.prompt-box {
    background: #14161b;
    border: 1px solid #2e3440;
    padding: 6px;
    border-radius: 4px;
    resize: vertical;
    min-height: calc(126px * var(--row-scale));
    max-height: calc(166px * var(--row-scale));
    overflow-y: auto;
    overflow-x: hidden;
    font-size: 12px;
    line-height: 1.5;
    color: #eceff4;
    white-space: pre-wrap !important;
    word-wrap: break-word !important;
    width: 100%;
    box-sizing: border-box;
    font-family: inherit;
}

/* 灯箱 */
#lightboxModal, #videoLightboxModal {
    display: none;
    position: fixed;
    z-index: 9999;
    left: 0;
    top: 0;
    width: 100%;
    height: 100%;
    background: rgba(0,0,0,0.88);
    justify-content: center;
    align-items: center;
}

#lightboxModal img, #videoLightboxModal video {
    max-width: 92%;
    max-height: 92%;
    border-radius: 6px;
    box-shadow: 0 0 20px rgba(0,0,0,0.9);
}
"""

# ==============================
# 4. JS 脚本
# ==============================
CUSTOM_JS = """
let wScale = parseFloat(localStorage.getItem('userColScale')) || 1.0;
let hScale = parseFloat(localStorage.getItem('userRowScale')) || 1.0;

function applyScales() {
    const root = document.documentElement;
    root.style.setProperty('--col-scale', wScale);
    root.style.setProperty('--row-scale', hScale);
}

window.adjustWidth = function(delta) {
    wScale = Math.max(0.6, Math.min(2.5, wScale + delta));
    applyScales();
};

/* 三个文本框都开了 resize: vertical，手动拖高会留下内联 height，
   它会盖住 CSS 里的 calc(... * var(--row-scale))，缩放就管不到这一格了。
   所以纵向缩放 / 恢复默认之前先把内联高度清掉，让 CSS 重新接管。 */
window.clearManualRowHeights = function() {
    document.querySelectorAll('.storyboard-table .area-title, .storyboard-table .dialogue-input, .storyboard-table .prompt-box')
        .forEach(function(el) { el.style.removeProperty('height'); });
};

window.adjustHeight = function(delta) {
    hScale = Math.max(0.6, Math.min(2.5, hScale + delta));
    clearManualRowHeights();
    applyScales();
};

window.saveLayoutScale = function() {
    localStorage.setItem('userColScale', wScale);
    localStorage.setItem('userRowScale', hScale);
    showToast('当前栏宽与行高已保存为默认！', 'success');
};

window.resetLayoutScale = function() {
    wScale = 1.0;
    hScale = 1.0;
    localStorage.removeItem('userColScale');
    localStorage.removeItem('userRowScale');
    clearManualRowHeights();
    applyScales();
    showToast('已恢复默认排版', 'success');
};

window.fitWidth = function() {
    const viewport = document.getElementById('storyboardViewport');
    if (!viewport) return;
    const containerWidth = viewport.clientWidth - 20;
    const table = viewport.querySelector('table');
    if (!table) return;
    const tableWidth = table.scrollWidth;
    wScale = containerWidth / tableWidth * wScale;
    wScale = Math.max(0.6, Math.min(1.5, wScale));
    applyScales();
};

window.showToast = function(msg, type = 'success') {
    let toast = document.getElementById('globalToast');
    if (!toast) {
        toast = document.createElement('div');
        toast.id = 'globalToast';
        toast.className = 'toast-box';
        document.body.appendChild(toast);
    }
    toast.className = 'toast-box toast-' + type;
    toast.textContent = msg;
    setTimeout(() => toast.classList.add('toast-show'), 10);
    clearTimeout(window._toastTimer);
    window._toastTimer = setTimeout(() => {
        toast.classList.remove('toast-show');
    }, 2500);
};

window.onRowClick = function(tr) {
    document.querySelectorAll('tr.storyboard-row').forEach(r => r.classList.remove('selected-row'));
    tr.classList.add('selected-row');
    const promptDiv = tr.querySelector('[data-col="prompt"]');
    const box = document.getElementById('masterPromptBox');
    if (promptDiv && box) {
        const text = promptDiv.value || '';
        box.innerHTML = "<b style='color:#ebcb8b'>[本镜提示词]：</b>" + (text ? text : "<i style='color:#8c9ba5'>无提示词内容</i>");
    }
};

window.openLightbox = function(url) {
    let modal = document.getElementById('lightboxModal');
    if (!modal) {
        modal = document.createElement('div');
        modal.id = 'lightboxModal';
        modal.innerHTML = '<img id="lightboxImg" src="">';
        modal.onclick = function() { modal.style.display = 'none'; };
        document.body.appendChild(modal);
    }
    document.getElementById('lightboxImg').src = url;
    modal.style.display = 'flex';
};

window.openVideoLightbox = function(src, btnEl) {
    event.stopPropagation();
    const wrap = btnEl.closest('.video-wrap');
    const smallVideo = wrap ? wrap.querySelector('.video-player') : null;

    let currentTime = 0;
    if (smallVideo) {
        currentTime = smallVideo.currentTime;
        smallVideo.pause();
    }

    let modal = document.getElementById('videoLightboxModal');
    if (!modal) {
        modal = document.createElement('div');
        modal.id = 'videoLightboxModal';
        modal.innerHTML = '<video id="lightboxVideo" controls preload="auto" onclick="event.stopPropagation()" playsinline></video>';
        modal.onclick = function() {
            const v = document.getElementById('lightboxVideo');
            if (v) v.pause();
            modal.style.display = 'none';
        };
        document.body.appendChild(modal);
    }

    const video = document.getElementById('lightboxVideo');
    video.src = src;
    video.currentTime = currentTime;
    modal.style.display = 'flex';
    video.play().catch(() => {});
};

window.applyColVisibility = function(selectedList) {
    const colMap = {
        "镜头号": "col-shot",
        "小标题": "col-title",
        "音频结果/镜头时长": "col-dur",
        "台词内容/台词角色": "col-dialogue",
        "图片资产": "col-img",
        "提示词": "col-prompt",
        "视频结果": "col-video"
    };
    Object.values(colMap).forEach(cls => {
        document.querySelectorAll('.' + cls).forEach(el => el.style.display = 'none');
    });
    selectedList.forEach(name => {
        const cls = colMap[name];
        if (cls) {
            document.querySelectorAll('.' + cls).forEach(el => el.style.display = '');
        }
    });
};

// Excel保存
window.saveToExcel = function() {
    const viewport = document.getElementById('storyboardViewport');
    const sheetName = viewport ? viewport.dataset.sheetName : '';

    if (!sheetName) {
        showToast('请先加载分镜表', 'error');
        return;
    }

    let rows = [];
    const trList = document.querySelectorAll('#storyboardTable tbody tr');
    
    for (let i = 0; i < trList.length; i++) {
        const tr = trList[i];
        rows.push({
            shot: tr.querySelector('[data-col="shot"]')?.value || "",
            title: tr.querySelector('[data-col="title"]')?.value || "",
            dur: tr.querySelector('[data-col="dur"]')?.value || "",
            dialogue: tr.querySelector('[data-col="dialogue"]')?.value || "",
            img: tr.querySelector('[data-col="img"]')?.value || "",
            prompt: tr.querySelector('[data-col="prompt"]')?.value || "",
            video: ""
        });
    }

    if (rows.length === 0) {
        showToast('表格数据为空', 'error');
        return;
    }

    fetch('/api/save_excel', {
        method: 'POST',
        headers: {'Content-Type': 'application/json'},
        body: JSON.stringify({
            sheet_name: sheetName,
            rows: rows
        })
    })
    .then(res => res.json())
    .then(res => {
        if (res.status === 'success') {
            showToast(res.msg, 'success');
        } else {
            showToast(res.msg, 'error');
        }
    })
    .catch(err => {
        showToast('保存失败：网络错误', 'error');
    });
};

window.pauseOtherVideos = function(currentVideo) {
    document.querySelectorAll('.video-player').forEach(v => {
        if (v !== currentVideo) v.pause();
    });
};

/* 全局快捷键（键盘上一排 a s d f g h，另加 r）
     a / s → 切「当前这张表」的上一个 / 下一个 sheet
             （表格 A = 上一表 / 下一表，表格 B = 上一场 / 下一场）
     d / f → 切 Tab 页面（d → 表格 A，f → 表格 B）
     g / h → 横向缩放；Shift+g / Shift+h → 纵向缩放；r → 恢复默认排版
   —— 只在「当前可见的那一部分」上生效，且焦点在输入框里时不拦截（方便打字）*/
document.addEventListener('keydown', function(e) {
    const el = document.activeElement;
    const activeTag = el ? el.tagName : '';
    if (activeTag === 'INPUT' || activeTag === 'TEXTAREA' || activeTag === 'SELECT') return;
    if (e.metaKey || e.ctrlKey || e.altKey) return;

    const key = e.key || '';
    const lower = key.toLowerCase();

    // a / s ：切 sheet（本表内翻页，等价于点「上一表 / 下一表」或「上一场 / 下一场」）
    if (lower === 'a' || lower === 's') {
        e.preventDefault();
        if (window.sbsStepSheet) sbsStepSheet(lower === 'a' ? -1 : 1);
        return;
    }

    // d / f ：切 Tab 页面（表格 A ↔ 表格 B）
    if (lower === 'd' || lower === 'f') {
        e.preventDefault();
        if (window.sbsSwitchTab) sbsSwitchTab(lower === 'd' ? 'A' : 'B');
        return;
    }

    // g / h ：横向缩放；按住 Shift（key 变大写）→ 纵向缩放
    if (lower === 'g' || lower === 'h') {
        // 判定纵向缩放必须用 shiftKey：Safari 下按 Shift+h，e.key 依旧是小写 'h'，
        // 靠大写字母判定会失效（Caps Lock 开着时还会误判，所以只认 shiftKey）。
        const vertical = !!e.shiftKey;
        const grow = (lower === 'h');
        const step = 0.1;
        e.preventDefault();
        const vpA = document.getElementById('assetViewport');
        const inPartA = !!(vpA && vpA.offsetParent !== null);
        if (inPartA) {
            if (vertical) aAdjustHeight(grow ? step : -step);
            else aAdjustWidth(grow ? step : -step);
        } else {
            if (vertical) adjustHeight(grow ? step : -step);
            else adjustWidth(grow ? step : -step);
        }
        return;
    }

    // r ：恢复默认排版（对当前所在的部分生效）
    if (lower === 'r') {
        e.preventDefault();
        const vpA = document.getElementById('assetViewport');
        const inPartA = !!(vpA && vpA.offsetParent !== null);
        if (inPartA) aResetScale(); else resetLayoutScale();
    }
});

document.addEventListener('wheel', function(e) {
    const viewport = document.getElementById('storyboardViewport');
    if (!viewport) return;
    if (!viewport.contains(e.target)) return;
    if (e.shiftKey) {
        e.preventDefault();
        viewport.scrollLeft += e.deltaY;
    }
}, { passive: false });

setTimeout(applyScales, 500);
"""

# 追加第一部分（资产表）的样式与脚本：只做字符串拼接，原有内容原封不动
CUSTOM_CSS = CUSTOM_CSS + ASSET_CSS
CUSTOM_JS = CUSTOM_JS + ASSET_JS

# ==============================
# 5.5 路径框右端的原生「浏览…」按钮
# ==============================
# Windows 版 Edge 不支持把文件拖成路径，所以给 5 个路径框各挂一个 📁 小按钮：
# 点一下由服务端弹系统选择框，选完自动回填（依然支持手动粘地址）。
# 按钮绝对定位在输入框右端，完全不占纵向空间。
PATH_PICKER_CSS = """
/* 路径框右端的「浏览…」小按钮：绝对定位覆盖在输入框右侧空白处 */
.sbs-input-wrap { position: relative; }
.sbs-input-wrap .sbs-pick-btn {
    position: absolute;
    right: 6px;
    top: 50%;
    transform: translateY(-50%);
    width: 26px;
    height: 26px;
    padding: 0;
    margin: 0;
    font-size: 13px;
    line-height: 24px;
    text-align: center;
    color: #e5e9f0;
    background: #343b4a;
    border: 1px solid #4a5266;
    border-radius: 5px;
    cursor: pointer;
    opacity: .85;
    z-index: 4;
}
.sbs-input-wrap .sbs-pick-btn:hover { opacity: 1; background: #46506a; }
.sbs-input-wrap .sbs-pick-btn:disabled { opacity: .45; cursor: default; }
/* 给按钮让位，免得长路径钻到按钮底下 */
.sbs-input-wrap input[data-testid="textbox"] { padding-right: 34px !important; }
"""

PATH_PICKER_JS = """
/* ===== 原生「浏览…」按钮：把系统选择框选到的路径直接回填到输入框 ===== */
(function () {
    var TARGETS = [
        { label: '表格 A (资产表)', kind: 'file', title: '选择表格 A（资产表）' },
        { label: '表格 B (分镜表)', kind: 'file', title: '选择表格 B（分镜表）' },
        { label: '资产目录', kind: 'dir', title: '选择资产目录' },
        { label: '音频目录', kind: 'dir', title: '选择音频目录' },
        { label: '视频目录', kind: 'dir', title: '选择视频目录' }
    ];

    function setInputValue(input, value) {
        input.value = value;
        input.dispatchEvent(new Event('input', { bubbles: true }));
        input.dispatchEvent(new Event('change', { bubbles: true }));
    }

    function mountOne(cfg) {
        var spans = document.querySelectorAll('span[data-testid="block-info"]');
        for (var i = 0; i < spans.length; i++) {
            var txt = (spans[i].textContent || '').trim();
            if (txt.indexOf(cfg.label) !== 0) continue;          // 只认以标签名开头的那个
            var blk = spans[i].closest('.block') || spans[i].parentElement;
            if (!blk) return false;
            var input = blk.querySelector('input[data-testid="textbox"]');
            if (!input) return false;
            var wrap = input.parentElement;
            if (!wrap) return false;
            if (wrap.querySelector('.sbs-pick-btn')) return true; // 已挂过，不重复
            var btn = document.createElement('button');
            btn.type = 'button';
            btn.className = 'sbs-pick-btn';
            btn.textContent = '📁';
            btn.title = cfg.kind === 'dir' ? '浏览文件夹…' : '浏览文件…';
            btn.addEventListener('click', function (ev) {
                ev.preventDefault();
                ev.stopPropagation();
                btn.disabled = true;
                btn.textContent = '⏳';
                fetch('/api/pick_path', {
                    method: 'POST',
                    headers: { 'Content-Type': 'application/json' },
                    body: JSON.stringify({ kind: cfg.kind, title: cfg.title })
                })
                    .then(function (r) { return r.json(); })
                    .then(function (j) {
                        if (j.status === 'ok' && j.path) {
                            setInputValue(input, j.path);
                            if (window.showToast) showToast('已填入：' + j.path, 'success');
                        } else if (j.status === 'error' && window.showToast) {
                            showToast(j.msg || '选择失败', 'error');
                        }
                    })
                    .catch(function () {
                        if (window.showToast) showToast('无法调用系统选择框，请手动粘贴路径', 'error');
                    })
                    .then(function () { btn.disabled = false; btn.textContent = '📁'; });
            });
            wrap.classList.add('sbs-input-wrap');
            wrap.appendChild(btn);
            return true;
        }
        return false;
    }

    function mountAll() { TARGETS.forEach(mountOne); }
    window.sbsMountPathPickers = mountAll;
    // Gradio 的 DOM 是异步渲染的：首屏补一次，之后低频轮询自愈（切 tab 重渲染也不会丢）
    setTimeout(mountAll, 800);
    setInterval(mountAll, 2500);
})();
"""

CUSTOM_CSS = CUSTOM_CSS + PATH_PICKER_CSS
CUSTOM_JS = CUSTOM_JS + PATH_PICKER_JS

# ==============================
# 5. 后端控制器
# ==============================
class StoryboardApp:
    def __init__(self):
        self.excel_b_path, self.asset_dir, self.video_dir, self.audio_dir = "", "", "", ""
        self.image_index = {}
        self.image_conflicts = []
        self.audio_index = {}
        self.audio_conflicts = []
        self.video_paths = []
        self.b_sheets, self.current_sheet_b = [], ""

    def load_project(self, path_a, path_b, asset_p, video_p, audio_p, current_sheet_val):
        save_config(path_a, path_b, asset_p, video_p, audio_p)
        self.excel_b_path = clean_path(path_b)
        self.asset_dir = clean_path(asset_p)
        self.video_dir = clean_path(video_p)
        self.audio_dir = clean_path(audio_p)

        self.image_index, self.image_conflicts = build_image_index(self.asset_dir)
        
        # 音频目录优先级
        target_audio_dir = self.audio_dir
        if not target_audio_dir:
            sub_audio = os.path.join(self.asset_dir, "Audio")
            if os.path.isdir(sub_audio):
                target_audio_dir = sub_audio
            else:
                target_audio_dir = self.asset_dir
        self.audio_index, self.audio_conflicts = build_audio_index(target_audio_dir)

        target_video_dir = self.video_dir if self.video_dir else self.asset_dir
        self.video_paths = build_video_list(target_video_dir)

        sheet_choices, default_sheet = [], None
        if os.path.exists(self.excel_b_path):
            try:
                wb = openpyxl.load_workbook(self.excel_b_path, read_only=True)
                # 严格按源文件里 sheet 的原有顺序（改名也照样排第一位）
                self.b_sheets = list(wb.sheetnames)
                sheet_choices = self.b_sheets
                default_sheet = current_sheet_val if current_sheet_val in sheet_choices else (self.b_sheets[0] if self.b_sheets else None)
                self.current_sheet_b = default_sheet
            except Exception as e:
                gr.Warning(f"读取分镜表失败: {e}")
        return gr.update(choices=sheet_choices, value=default_sheet)

    def get_conflict_html(self):
        items = []
        if self.image_conflicts:
            items.append(f"<span class='conflict-item'>📷 图片冲突：{'<br>'.join(self.image_conflicts)}</span>")
        if self.audio_conflicts:
            items.append(f"<span class='conflict-item'>🔊 音频冲突：{'<br>'.join(self.audio_conflicts)}</span>")
        return f"<div class='conflict-panel'>{''.join(items)}</div>"

    def load_sheet_b_content(self, sheet_name):
        if not self.excel_b_path or not os.path.exists(self.excel_b_path) or not sheet_name:
            return "<div class='empty-tip'>请先加载分镜表</div>"
        self.current_sheet_b = sheet_name
        df = pd.read_excel(self.excel_b_path, sheet_name=sheet_name)
        return render_table_b_html(df, self.image_index, self.audio_index, self.video_paths, sheet_name)

    def step_sheet(self, direction, current_sheet):
        if not self.b_sheets or current_sheet not in self.b_sheets:
            return current_sheet, self.load_sheet_b_content(current_sheet)
        new_idx = self.b_sheets.index(current_sheet) + direction
        if 0 <= new_idx < len(self.b_sheets):
            ts = self.b_sheets[new_idx]
            return ts, self.load_sheet_b_content(ts)
        return current_sheet, self.load_sheet_b_content(current_sheet)

app_core = StoryboardApp()
init_conf = load_config()

# ==============================
# 6. UI 页面构建
# ==============================
with gr.Blocks(title="StoryBoardStudio 分镜工作台") as demo:
    with gr.Row():
        txt_path_a = gr.Textbox(label="表格 A (资产表)", value=init_conf.get("path_a", ""), lines=1, max_lines=1, scale=2)
        txt_path_b = gr.Textbox(label="表格 B (分镜表)", value=init_conf.get("path_b", ""), lines=1, max_lines=1, scale=2)
        txt_asset_dir = gr.Textbox(label="资产目录", value=init_conf.get("asset_dir", ""), lines=1, max_lines=1, scale=2)
        txt_audio_dir = gr.Textbox(label="音频目录", value=init_conf.get("audio_dir", ""), lines=1, max_lines=1, scale=1, placeholder="不填默认资产目录/Audio")
        txt_video_dir = gr.Textbox(label="视频目录", value=init_conf.get("video_dir", ""), lines=1, max_lines=1, scale=1, placeholder="不填默认同资产目录")
        btn_load_project = gr.Button("🚀 加载/刷新工程", variant="primary", scale=0, size="sm", min_width=104)

    with gr.Tabs():

        # ===== 第一部分：影视资产表（X/Y/Z 预览 + 听声）=====
        build_asset_tab(txt_path_a, txt_asset_dir, txt_audio_dir, btn_load_project)
        with gr.TabItem("📋 表格 B · 影视分镜表"):
            # 冲突提示区
            conflict_panel = gr.HTML('<div class="conflict-panel"></div>')

            # 第一行：场次 + 隐藏列 + 上下场
            with gr.Row(elem_classes="top-toolbar-row"):
                dd_sheets_b = gr.Dropdown(label="当前场次", choices=[], interactive=True, scale=2)
                with gr.Column(scale=3, elem_classes="col-checkbox-wrap"):
                    chk_cols = gr.CheckboxGroup(
                        choices=["镜头号", "小标题", "音频结果/镜头时长", "台词内容/台词角色", "图片资产", "提示词", "视频结果"],
                        value=["镜头号", "小标题", "音频结果/镜头时长", "台词内容/台词角色", "图片资产", "提示词", "视频结果"],
                        label="✓ 隐藏/显示列",
                        scale=0
                    )
                btn_prev_sheet = gr.Button("◀ 上一场", scale=0, min_width=60)
                btn_next_sheet = gr.Button("下一场 ▶", scale=0, min_width=60)

            # 第二行：视图微调
            with gr.Row(elem_classes="view-toolbar-row"):
                gr.Markdown("**📐 视图微调**", scale=0)
                btn_w_dec = gr.Button("◀ 栏宽收窄", scale=1)
                btn_w_inc = gr.Button("▶ 栏宽拉宽", scale=1)
                btn_h_dec = gr.Button("▲ 行高收紧", scale=1)
                btn_h_inc = gr.Button("▼ 行高拉开", scale=1)
                btn_fit_width = gr.Button("↔ 适配宽度", scale=1)
                btn_scale_save = gr.Button("📌 记住排版", scale=1)
                btn_scale_reset = gr.Button("🔄 恢复默认", scale=1)

            # 第三行：主提示词 + 保存按钮
            with gr.Row(elem_classes="prompt-save-row"):
                gr.HTML("""<div id="masterPromptBox" class="master-prompt-preview"></div>""")
                btn_save_to_excel = gr.Button("💾 同步并保存到 Excel", variant="primary", scale=0, elem_classes="save-btn")

            html_storyboard_view = gr.HTML("<div class='empty-tip'>等待加载工程数据...</div>")

    # 事件绑定
    btn_w_dec.click(None, None, None, js="adjustWidth(-0.1)")
    btn_w_inc.click(None, None, None, js="adjustWidth(0.1)")
    btn_h_dec.click(None, None, None, js="adjustHeight(-0.1)")
    btn_h_inc.click(None, None, None, js="adjustHeight(0.1)")
    btn_fit_width.click(None, None, None, js="fitWidth()")
    btn_scale_save.click(None, None, None, js="saveLayoutScale()")
    btn_scale_reset.click(None, None, None, js="resetLayoutScale()")

    chk_cols.change(None, inputs=[chk_cols], js="(cols) => { applyColVisibility(cols); }")

    def load_and_refresh(path_a, path_b, asset_p, video_p, audio_p, cur_sheet):
        sheet_update = app_core.load_project(path_a, path_b, asset_p, video_p, audio_p, cur_sheet)
        conflict_html = app_core.get_conflict_html()
        return sheet_update, conflict_html

    btn_load_project.click(
        fn=load_and_refresh,
        inputs=[txt_path_a, txt_path_b, txt_asset_dir, txt_video_dir, txt_audio_dir, dd_sheets_b],
        outputs=[dd_sheets_b, conflict_panel]
    )

    dd_sheets_b.change(fn=app_core.load_sheet_b_content, inputs=[dd_sheets_b], outputs=[html_storyboard_view])
    btn_prev_sheet.click(fn=lambda cur: app_core.step_sheet(-1, cur), inputs=[dd_sheets_b], outputs=[dd_sheets_b, html_storyboard_view])
    btn_next_sheet.click(fn=lambda cur: app_core.step_sheet(1, cur), inputs=[dd_sheets_b], outputs=[dd_sheets_b, html_storyboard_view])

    btn_save_to_excel.click(None, None, None, js="saveToExcel()")

# ==============================
# 7. FastAPI 接口
# ==============================
fastapi_app = FastAPI()
# 第一部分（资产表）的保存接口 —— 实现放在 asset_tab.py，这里只注册
register_asset_api(fastapi_app)


# ---- 原生「浏览…」选择框 ----------------------------------------------------
# SBS 跑在本机，浏览器和服务端是同一台机器，所以文件/文件夹选择框可以直接由
# 服务端弹出（Windows 版 Edge 不支持把文件拖成路径，靠它就不用手敲地址了）。
def _pick_path_native(kind: str, title: str) -> str:
    """弹出系统原生选择框，返回绝对路径。

    用户取消 → 返回 ""；系统不支持或调用出错 → 抛 RuntimeError（前端会红字提示）。
    """
    title = (title or ("选择文件夹" if kind == "dir" else "选择文件"))
    # 标题只作提示，清掉会破坏脚本字符串字面量的字符
    title = title.replace('"', " ").replace("\\", " ").replace("'", " ")

    # ---------- macOS ----------
    if sys.platform == "darwin":
        verb = "choose folder" if kind == "dir" else "choose file"
        script = ("try\n"
                  f'    set t to {verb} with prompt "{title}"\n'
                  "    return POSIX path of t\n"
                  "on error number -128\n"          # -128 = 用户点了取消
                  '    return ""\n'
                  "end try")
        r = subprocess.run(["osascript", "-e", script], capture_output=True, timeout=600)
        err = (r.stderr or b"").decode("utf-8", "replace").strip()
        if r.returncode != 0 and err and "cancel" not in err.lower():
            raise RuntimeError(err.splitlines()[-1] if err else "osascript 调用失败")
        return _trim_sep((r.stdout or b"").decode("utf-8", "replace").strip())

    # ---------- Windows ----------
    if os.name == "nt":
        exe = shutil.which("powershell") or shutil.which("pwsh")
        if not exe:
            raise RuntimeError("未找到 powershell，请手动粘贴路径")
        if kind == "dir":
            pick = ("$d = New-Object System.Windows.Forms.FolderBrowserDialog;"
                    f"$d.Description = '{title}';"
                    "$d.ShowNewFolderButton = $true;"
                    "if ($d.ShowDialog() -eq [System.Windows.Forms.DialogResult]::OK)"
                    " { [Console]::Out.Write($d.SelectedPath) }")
        else:
            pick = ("$d = New-Object System.Windows.Forms.OpenFileDialog;"
                    f"$d.Title = '{title}';"
                    "$d.Filter = 'All files (*.*)|*.*';"
                    "$d.CheckFileExists = $true; $d.RestoreDirectory = $true;"
                    "if ($d.ShowDialog() -eq [System.Windows.Forms.DialogResult]::OK)"
                    " { [Console]::Out.Write($d.FileName) }")
        # [Console]::OutputEncoding 设成 UTF-8，中文路径才不会在回读时变乱码
        ps = ("Add-Type -AssemblyName System.Windows.Forms;"
              "[Console]::OutputEncoding = [System.Text.Encoding]::UTF8;" + pick)
        r = subprocess.run(
            [exe, "-NoProfile", "-STA", "-ExecutionPolicy", "Bypass", "-Command", ps],
            capture_output=True, timeout=600,
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),   # 不闪黑框
        )
        return _trim_sep((r.stdout or b"").decode("utf-8", "replace").lstrip("\ufeff").strip())

    # ---------- Linux ----------
    if shutil.which("zenity"):
        args = ["zenity", "--file-selection", f"--title={title}"]
        if kind == "dir":
            args.append("--directory")
        r = subprocess.run(args, capture_output=True, timeout=600)
        if r.returncode != 0:                       # 取消
            return ""
        return _trim_sep((r.stdout or b"").decode("utf-8", "replace").strip())
    raise RuntimeError("当前系统没有 zenity，请手动粘贴路径")


class PickPathRequest(BaseModel):
    kind: str = "file"          # "file" 选文件 / "dir" 选文件夹
    title: str = ""


@fastapi_app.post("/api/pick_path")
async def pick_path_api(req: PickPathRequest):
    """弹出系统原生选择框。取消返回 status=cancel（前端静默处理）。"""
    try:
        # 选择框会一直阻塞到用户选完，扔到线程里跑，免得把整个服务卡住
        picked = await asyncio.to_thread(_pick_path_native, req.kind, req.title)
    except Exception as e:                                   # noqa: BLE001
        return {"status": "error", "msg": f"调用系统选择框失败：{e}"}
    if not picked:
        return {"status": "cancel", "msg": "已取消选择"}
    return {"status": "ok", "path": picked}


class SaveExcelRequest(BaseModel):
    sheet_name: str
    rows: list

@fastapi_app.post("/api/save_excel")
def save_excel_api(req: SaveExcelRequest):
    try:
        sheet_name = req.sheet_name
        rows = req.rows

        if not app_core.excel_b_path or not os.path.exists(app_core.excel_b_path):
            return {"status": "error", "msg": "保存失败：Excel 文件不存在"}
        if not sheet_name or not rows:
            return {"status": "error", "msg": "保存失败：数据不完整"}

        # 备份到「资产目录/excel_backup/」（资产目录没填时退回 Excel 同目录 _backups/）
        backup_path = backup_excel(app_core.excel_b_path, app_core.asset_dir)

        wb = openpyxl.load_workbook(app_core.excel_b_path)
        ws = wb[sheet_name]
        for idx, row_data in enumerate(rows):
            r = idx + 2
            ws.cell(row=r, column=1, value=row_data.get('shot', ''))
            ws.cell(row=r, column=2, value=row_data.get('title', ''))
            ws.cell(row=r, column=3, value=row_data.get('dur', ''))
            ws.cell(row=r, column=4, value=row_data.get('dialogue', ''))
            ws.cell(row=r, column=5, value=row_data.get('img', ''))
            ws.cell(row=r, column=6, value=row_data.get('prompt', ''))
            ws.cell(row=r, column=7, value='')
        wb.save(app_core.excel_b_path)
        filename = os.path.basename(app_core.excel_b_path)
        if backup_path:
            return {"status": "success",
                    "msg": f"✅ 已保存 {filename}｜备份：{os.path.basename(backup_path)}"}
        return {"status": "success", "msg": f"✅ 已保存到 {filename}"}

    except PermissionError:
        return {"status": "error", "msg": "❌ 保存失败：Excel 文件正在被占用，请关闭 Excel/WPS 后重试"}
    except Exception as e:
        return {"status": "error", "msg": f"❌ 保存失败: {str(e)}"}

@fastapi_app.get("/api/local_media")
def serve_local_media(filepath: str):
    path = urllib.parse.unquote(filepath)
    if not os.path.exists(path):
        return Response(content=b"File not found", status_code=404)

    ext = os.path.splitext(path)[1].lower()

    if ext in {'.jpg', '.jpeg', '.png', '.webp', '.bmp', '.tif', '.tiff'}:
        try:
            with Image.open(path) as img:
                try:
                    img = ImageOps.exif_transpose(img)
                except Exception:
                    pass
                if img.mode not in ('RGB', 'L'):
                    img = img.convert('RGB')
                buf = io.BytesIO()
                img.save(buf, format="JPEG", quality=90, optimize=True)
                buf.seek(0)
                return Response(content=buf.getvalue(), media_type="image/jpeg")
        except Exception:
            pass

    mime_type, _ = mimetypes.guess_type(path)
    return FileResponse(path, media_type=mime_type or "application/octet-stream")

app = gr.mount_gradio_app(fastapi_app, demo, path="/", css=CUSTOM_CSS, js=CUSTOM_JS)

def _port_in_use(port: int, host: str = "0.0.0.0") -> bool:
    """启动前自检：端口被占，说明同款程序已经在跑了，不必再起一个。"""
    import socket
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
        try:
            s.bind((host, port))
            return False
        except OSError:
            return True


if __name__ == "__main__":
    import uvicorn

    PORT = 7861
    if _port_in_use(PORT):
        # 这不是故障，只是「重复启动」。给人话，别丢 errno 让人心慌。
        print("\n" + "=" * 58)
        print("  ⚠️  端口 7861 已被占用：StoryBoardStudio 应该已经在运行了")
        print("     直接打开浏览器 →  http://localhost:7861")
        print("     确实要重启的话，先在原窗口按 Ctrl+C，或执行：")
        print("         pkill -f StoryBoardStudio.py")
        print("=" * 58 + "\n")
        raise SystemExit(1)

    print(f"▶ StoryBoardStudio 已启动 → http://localhost:{PORT}")
    uvicorn.run(app, host="0.0.0.0", port=PORT)
